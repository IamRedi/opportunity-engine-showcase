"""CV parsing (Faza 4D) — ekstraktim teksti + parse deterministik me dëshmi.

Dy faza (spec 4D):
1. Ekstraktim teksti nga PDF/DOCX/TXT/MD (metoda e regjistruar; pa OCR).
2. Parse deterministic në sugjerime me dëshmi (rreshti burim + seksion) —
   pending derisa përdoruesi i aprovon; kurrë s'prek ProfileContext direkt.

Parseri është pa AI (4B-adapteri mund të shtohet më vonë me të njëjtat garda).
"""
import re
from pathlib import Path

PARSER_VERSION = "cv.3"

_TECH_VOCAB = ("python", "javascript", "typescript", "react", "vue", "angular", "node",
               "rust", "golang", "java", "kotlin", "swift", "php", "ruby", "rails",
               "django", "fastapi", "laravel", ".net", "c#", "c++", "sql", "postgresql",
               "mysql", "mongodb", "aws", "azure", "gcp", "docker", "kubernetes",
               "terraform", "machine learning", "devops", "flutter", "react native",
               "wordpress", "shopify", "next.js", "graphql", "redis")

_LANG_HINTS = ("gjermanisht", "german", "deutsch", "anglisht", "english", "englisch",
               "shqip", "albanian", "italian", "italiano", "französisch", "french",
               "francais", "spanish", "spanisch", "turkish", "turkish")


# ------------------------------------------------------------ ekstraktimi

def extract_text(path: Path) -> dict:
    """Kthen {text, method, status, error}. Nuk modifikon skedarin origjinal."""
    suffix = path.suffix.lower()
    try:
        import storage
        raw = storage.read_reference(path)
    except OSError as e:
        return {"text": "", "method": suffix, "status": "failed", "error": str(e)}

    if suffix == ".pdf" or raw[:5] == b"%PDF-":
        return _extract_pdf(raw, path)
    if suffix == ".docx":
        return _extract_docx(raw, path)
    if suffix in (".txt", ".md"):
        text = raw.decode('utf-8', errors='replace')
        if not text.strip():
            return {'text': '', 'method': 'utf8-text', 'status': 'no_text', 'error': 'CV-ja është bosh.'}
        return {"text": text,
                "method": "utf8-text", "status": "ok", "error": ""}
    return {"text": "", "method": suffix, "status": "unsupported",
            "error": f"format i pambështetur: {suffix}"}


def _extract_pdf(raw: bytes, path: Path) -> dict:
    try:
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > 100:
            return {'text': '', 'method': 'pypdf', 'status': 'failed', 'error': 'PDF me më shumë se 100 faqe; ngarko vetëm CV-në.'}
        pages = []
        layouts = []
        for i, page in enumerate(reader.pages[:100]):
            t = page.extract_text() or ""
            if t.strip():
                pages.append(f"[page {i + 1}]\n{t}")
                try:
                    layout = page.extract_text(extraction_mode='layout') or ''
                    layouts.append(f"[visual rows, page {i + 1}]\n{layout}")
                except Exception:
                    pass
        text = "\n".join(pages).strip()
        if not text:
            return {"text": "", "method": "pypdf", "status": "no_text",
                    "error": "PDF pa tekst të ekstraktueshëm (skanuar? duhet OCR)"}
        return {"text": text, "layout_text": '\n'.join(layouts), "method": "pypdf", "status": "ok", "error": ""}
    except Exception as e:
        return {"text": "", "method": "pypdf", "status": "failed",
                "error": f"PDF i korruptuar: {e}"}


def _extract_docx(raw: bytes, path: Path) -> dict:
    try:
        import io
        import docx
        doc = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if not text:
            return {"text": "", "method": "python-docx", "status": "no_text",
                    "error": "DOCX pa tekst"}
        return {"text": text, "method": "python-docx", "status": "ok", "error": ""}
    except Exception as e:
        return {"text": "", "method": "python-docx", "status": "failed",
                "error": f"DOCX i korruptuar: {e}"}


# ------------------------------------------------- parse deterministic

def _norm(v: str) -> str:
    return re.sub(r"\s+", " ", (v or "").lower()).strip()


_SECTION_PATTERNS = {
    "skills": ("skills", "technical skills", "kompetenzen", "kenntnisse", "fähigkeiten", "aftësi", "aftësitë", "tech stack", "technologies"),
    "languages": ("languages", "language", "sprachen", "gjuhë", "gjuhët"),
    "experience": ("experience", "work experience", "professional experience", "berufserfahrung", "erfahrung", "punësim", "eksperienca", "përvoja", "përvojë pune"),
    "education": ("education", "ausbildung", "bildung", "studium", "arsimimi", "edukimi", "studime"),
    "certifications": ("certifications", "certificates", "zertifikate", "zertifizierungen", "certifikime", "certifikimet"),
    "portfolio": ("portfolio", "portofol"),
    "projects": ("projects", "projekte", "projektet"),
    "summary": ("summary", "professional summary", "profile", "profil", "profil professionnel", "über mich", "rreth meje", "përmbledhje", "përmbledhja", "objective"),
    "contact": ('contact', 'contacts', 'kontakt', 'kontakte', 'kontakti', 'personal details', 'persönliche daten'),
}


def _line_section(line: str, current: str) -> str:
    low = line.lower().strip().lstrip('# ').rstrip(":").strip()
    for section, hints in _SECTION_PATTERNS.items():
        if any(low == h or low.startswith(h + ':') for h in hints):
            return section
    return current


def _evidence_line(lines: list[str], idx: int) -> str:
    """Dëshmi: rreshti burim + rreshtat para/s pas si kontekst i shkurtër."""
    ctx = lines[max(0, idx - 1):idx + 2]
    return " ⏎ ".join(c.strip() for c in ctx if c.strip())[:300]


def parse_suggestions(text: str) -> list[dict]:
    """Parse deterministic → lista e sugjerimeve (pending) me dëshmi.
    Deterministik: i njëjti tekst → e njëjta listë."""
    suggestions: dict = {}
    lines = text.splitlines()

    def add(section, field, proposed, evidence, confidence=0.9):
        proposed = (proposed or "").strip()
        normalized = _norm(proposed)
        if not proposed or len(normalized) < 2:
            return
        key = (section, field, normalized)
        if key not in suggestions:
            suggestions[key] = {
                "section": section, "field": field, "proposed_value": proposed,
                "normalized_value": normalized, "evidence": evidence[:300],
                "confidence": confidence, "uncertain": confidence < 0.8,
            }

    # 1) seksionet me rreshta (header-i ndryshon seksionin aktual)
    current = ""
    summary_lines = []
    first_section_row = False
    for idx, raw in enumerate(lines):
        line = raw.strip().lstrip('# ')
        if not line:
            continue
        heading = _line_section(line, "")
        if heading:
            current = heading
            first_section_row = True
            line = line.split(':', 1)[1].strip() if ':' in line else ''
            if not line:
                continue
        elif re.match(r'^(?:contact|kontakt|kontakte|contacts)\s*:', line, re.I):
            current = ''
        elif line.endswith(':'):
            current = ''
            continue

        # aftësi brenda seksionit skills: tokene të ndara me presje/pikëpresje
        if current == "skills":
            for token in re.split(r"[,;•·|]+", line):
                token = token.strip(" -–•·")
                if 2 <= len(token) <= 60 and not token.lower().startswith("skill"):
                    add("skills", "name", token, _evidence_line(lines, idx), 0.7 if token.endswith('.') and len(token.split()) > 4 else 0.9)

        # gjuhët: "Emri – nivel" ose "Emri (B2)"
        if current == "languages":
            m = re.match(r"([A-Za-zËËÇçÄäÖöÜüß ]{2,20})\s*[–—|(-]\s*([A-C]\d{0,2}|native|amtare|amtlich)", line, re.IGNORECASE)
            if m:
                add("languages", "name", f"{m.group(1).strip()} – {m.group(2).strip()}",
                    _evidence_line(lines, idx), 0.9)

        # eksperienca: "Titull @ Organizata" ose "Titull | Organizata | periudhë"
        if current == "experience":
            m = re.match(r"(.+?)\s*(@|\||at)\s*(.+)$", line)
            if m and 3 <= len(m.group(1)) <= 80:
                add("experience", "title",
                    f"{m.group(1).strip()} @ {m.group(3).strip()}",
                    _evidence_line(lines, idx), 0.8)

        # edukimi: rreshtat brenda seksionit education
        if current == "education" and 3 <= len(line) <= 120:
            add("education", "title", line, _evidence_line(lines, idx), 0.8 if first_section_row else 0.6)

        # certifikimet
        if current == "certifications" and 3 <= len(line) <= 120:
            add("certifications", "name", line, _evidence_line(lines, idx), 0.85 if first_section_row else 0.6)

        # përmbledhje: teksti i plotë i seksionit (merge vetëm nëse bosh)
        if current == "summary" and len(line) > 40:
            summary_lines.append(line)
        first_section_row = False

    if summary_lines:
        summary = '\n'.join(summary_lines)[:3000]
        add('summary', 'text', summary, summary[:300], 0.85)

    # Explicit labels are reliable in all three interface languages.
    labels = {'name': 'name|full name|emri|emër|vorname und nachname',
              'title': 'title|job title|professional title|titulli profesional|berufsbezeichnung',
              'location': 'location|address|adresse|wohnort|vendndodhja|adresa'}
    for raw in lines:
        for field, aliases in labels.items():
            m = re.match(rf'^\s*(?:{aliases})\s*:\s*(.+)$', raw, re.I)
            if m:
                add('contact', field, m.group(1), raw.strip(), 0.95)

    # 2) vocabulary-skan mbi gjithë tekstin (aftësi teknike, confidence më i ulët)
    low = text.lower()
    for tech in _TECH_VOCAB:
        for idx, line in enumerate(lines):
            if tech in line.lower():
                m = re.search(rf"\b{re.escape(tech)}\b", line.lower())
                if m:
                    add("skills", "name", tech, _evidence_line(lines, idx), 0.6)
                    break  # një dëshmi për aftësi mjafton

    # 3) gjuhët e deklaruara me nivele në çdo vend
    for idx, raw in enumerate(lines):
        m = re.search(r"(gjermanisht|german|deutsch|anglisht|english|englisch|shqip|italian|"
                      r"französisch|french|spanish|spanisch|turkish)\s*[–—:-]?\s*"
                      r"([ABC]\d{1,2}|native|amtare|mother tongue)", raw, re.IGNORECASE)
        if m:
            add("languages", "name", f"{m.group(1).capitalize()} – {m.group(2).upper()}",
                _evidence_line(lines, idx), 0.9)

    # 4) URL-t publike → portfolio
    for idx, raw in enumerate(lines):
        for m in re.finditer(r"https?://[^\s,;)]+", raw):
            url = m.group(0).rstrip(".")
            if any(s in url for s in ("github", "gitlab", "portfolio", "linkedin.com/in")):
                add("portfolio", "title", url, _evidence_line(lines, idx), 0.85)

    # 5) kontaktet: email / telefon (për metadata)
    for m in re.finditer(r"[\w.+-]+@[\w-]+\.[\w.]+", text):
        add("contact", "email", m.group(0), "nga teksti i CV-së", 0.95)
        break
    for raw in lines:
        # Avoid treating an employment date range as a telephone number.
        m = re.search(r'(?:\+\d[\d ()/-]{6,22}\d)|(?:(?:phone|telefon|tel|mobile)\s*:\s*(\d[\d ()/-]{6,22}\d))', raw, re.I)
        if m:
            add('contact', 'phone', m.group(1) or m.group(0), raw.strip(), 0.9)
            break

    return list(suggestions.values())
